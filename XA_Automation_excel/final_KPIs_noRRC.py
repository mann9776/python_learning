import datetime

import openpyxl
from openpyxl import Workbook
from openpyxl import load_workbook
from openpyxl.styles import Alignment
import pandas as pd
from datetime import datetime

KPIs_header = ['DL MCS QPSK', 'DL MCS 16QAM', 'DL MCS 64QAM', 'DL MCS 256QAM', 'Data_QoS_Flow_5QI9_Setup_Success_Rate', 'UE_Call_Drop_Rate', 'VoNR Call Setup (5QI1)', 'VoNR Drop Rate (5QI1)', 'VoNR Call Setup Time (5QI1)', 'Intra_NR_Overall_Execution_Success_Rate', 'NR_RAN_Cell_Availability_Rate', 'PDUSessionEstSucessRate_Each_Slice', 'DL BLER', 'UL BLER', 'DL PRB %', 'UL PRB %', 'PHY-DL user TPUT', 'PHY-UL user TPUT', 'PAGING DISCARD at CU', 'PAGING DISCARD at DU', 'DynScheduledUeDlMax', 'DynScheduledUeUlMax', 'FiveQI_Abnormal_release_percentage', 'Average_Active_UE_DL_All_QCI', 'Average_Active_UE_UL_All_QCI']
KPIs_header1 = ('DL MCS QPSK', 'DL MCS 16QAM', 'DL MCS 64QAM', 'DL MCS 256QAM', 'Data_QoS_Flow_5QI9_Setup_Success_Rate', 'UE_Call_Drop_Rate', 'VoNR Call Setup (5QI1)', 'VoNR Drop Rate (5QI1)', 'VoNR Call Setup Time (5QI1)', 'Intra_NR_Overall_Execution_Success_Rate', 'NR_RAN_Cell_Availability_Rate', 'PDUSessionEstSucessRate_Each_Slice', 'PAGING DISCARD at CU', 'PAGING DISCARD at DU', 'FiveQI_Abnormal_release_percentage')

def load_sheet(path1):
    List_of_Dict = []
    wb = load_workbook(filename=path1)
    for KPI in KPIs_header:
        for sheet in wb:
            Dict = {}
            #print(sheet["A1"].value)
            for i in range(1, sheet.max_column+1):
                if sheet.cell(row=1, column=i).value == KPI:
                    if KPI == 'RRE_FAILURE_RATE (DROP RATE)':
                        for i1 in range(1, sheet.max_column+1):
                            if sheet.cell(row=1, column=i1).value == "RRC_ReEstabAtt_Sum":
                                for i2 in range(1, sheet.max_column+1):
                                    if sheet.cell(row=1, column=i2).value == "RRC_ReEstabSuccWithUeContext_Sum":
                                        for i3 in range(1, sheet.max_column+1):
                                            if sheet.cell(row=1, column=i3).value == "RRE_attempt_in_source_cell":
                                                for j in range(2, sheet.max_row+1):
                                                    Dict[sheet.cell(row=j, column=1).value] = [sheet.cell(row=j, column=i).value, sheet.cell(row=j, column=i1).value, sheet.cell(row=j, column=i2).value, sheet.cell(row=j, column=i3).value]
                                                List_of_Dict.append(Dict)
                                                break
                                        break
                                break
                        break
                    elif KPI == 'RRC_Reestablishment_Fallback_Success_Rate':
                        for i1 in range(1, sheet.max_column+1):
                            if sheet.cell(row=1, column=i1).value == "RRC_ReEstabFallbackToSetupAtt":
                                for i2 in range(1, sheet.max_column+1):
                                    if sheet.cell(row=1, column=i2).value == "RRC_ReEstabSuccWithoutUeContext":
                                        for j in range(2, sheet.max_row+1):
                                            Dict[sheet.cell(row=j, column=1).value] = [sheet.cell(row=j, column=i).value, sheet.cell(row=j, column=i1).value, sheet.cell(row=j, column=i2).value]
                                        List_of_Dict.append(Dict)
                                        break
                                break
                        break
                    else :
                        for j in range(2, sheet.max_row+1):
                        #print(sheet.cell(row=j, column=i).value)
                            Dict[sheet.cell(row=j, column=1).value] = sheet.cell(row=j, column=i).value
                        List_of_Dict.append(Dict)
                        break

    return List_of_Dict


def create_sheet(path):
    wb = Workbook()
    sheet = wb.active
    #print(list_value[0])
    #print(type(list_value))
    for KPI in KPIs_header:
        list_value = list(list_of_dict[KPIs_header.index(KPI)].items())
        if KPI == 'RRE_FAILURE_RATE (DROP RATE)':
            sheet.cell(row=KPIs_header.index(KPI) + 1, column=1).value = KPI+'\nRRC_ReEstabAtt_Sum'+'\nRRC_ReEstabSuccWithUeContext_Sum'+'\nRRE_attempt_in_source_cell'
        elif KPI == 'RRC_Reestablishment_Fallback_Success_Rate':
            sheet.cell(row=KPIs_header.index(KPI) + 1, column=1).value = KPI+'\nRRC_ReEstabFallbackToSetupAtt'+'\nRRC_ReEstabSuccWithoutUeContext'
        else:
            sheet.cell(row=KPIs_header.index(KPI)+1, column=1).value = KPI
        x = len(list_value)
        cell_value = ''
        for i in range(x):
            tup = list_value[i]
            kpi_value = ''
            for item in tup:
                if KPI in KPIs_header1:
                    kpi_value = str(item)
                else:
                    kpi_value = f'{kpi_value}   {item}'
            cell_value = cell_value+kpi_value+"\n"
        sheet.cell(row=KPIs_header.index(KPI)+1, column=2).value = cell_value
        sheet.cell(row=KPIs_header.index(KPI)+1, column=2).alignment = Alignment(wrapText=True)
    #sheet.merge_cells(start_row=1, end_row=x , start_column=2, end_column=2)
    sheet.column_dimensions['A'].width = 38
    sheet.column_dimensions['B'].width = 47
    wb.save(path)


if __name__ == "__main__":
    list_of_dict = load_sheet("C:/KPI_Tool/18-Cell/test.xlsx")
    print(list_of_dict)
    create_sheet("C:/KPI_Tool/18-Cell/KPIs_output_"+str(datetime.now().strftime("%Y-%m-%d %H-%M-%S"))+".xlsx")