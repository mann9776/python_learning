import openpyxl
from openpyxl import Workbook
from openpyxl import load_workbook
from openpyxl.styles import Alignment
import pandas as pd
#workbook = openpyxl.Workbook()
#sheet = workbook.active
#data = [["Name", "Age", "City"], ["John", 28, "New York"], ["Alice", 24, "SanFrancisco"], ["Bob", 32, "Los Angeles"]]
#for row in data:
#    sheet.append(row)
#workbook.save("C:/Users/singhman/Downloads/learn_PyExcel/PyExcel.xlsx")
#print("Excel File Created Successfully")



def create_workbook(path):
    wb = Workbook()
    sheet = wb.active
    #sheet.title = "Hello"
    wb.create_sheet()
    print(wb.sheetnames)
    sheet2 = wb.create_sheet(index=1, title="World")
    print(wb.sheetnames)
    sheet["A1"] = "Hello"
    sheet["B1"] = "from"
    sheet["C1"] = "OpenPyXL"
    sheet["A2"] = "row 2"
    sheet["A3"] = "row 3"
    sheet["B4"] = "row 4"
    #insert a column before column A
    #sheet.insert_cols(idx=1)

    #Delete column A
#    sheet.delete_cols(idx=1)

    #insert 2 rows starting from second row
    #sheet.insert_rows(idx=2, amount=2)

    #Delete 2 rows starting on the second row
#    sheet.delete_rows(idx=2, amount=2)

    del wb["World"]
    print(wb.sheetnames)
    wb.save(path)

def edit(path, data):
    wb = load_workbook(filename=path)
    sheet = wb.active
    for cell in data:
        current_value = sheet[cell].value
        sheet[cell] = data[cell]
        print(f'Changing {cell} from {current_value} to {data[cell]}')
    wb.save(path)

def merged_cells(path, value):
    wb = Workbook()
    sheet = wb.active
    sheet.merge_cells("A2:E2")
    top_left_cell = sheet["A2"]
    top_left_cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet["A2"] = value
    wb.save(path)

def folding(path, rows=None, cols=None, hidden=True):
    wb = Workbook()
    sheet = wb.active
    if rows:
        begin_row, end_row = rows
        sheet.row_dimensions.group(begin_row, end_row, hidden=hidden)
    if cols:
        begin_col, end_col = cols
        sheet.column_dimensions.group(begin_col, end_col, hidden=hidden)
    wb.save(path)

def freeze(path, row_to_freeze):
    wb = Workbook()
    sheet = wb.active
    sheet.title = "Freeze"
    sheet.freeze_panes = row_to_freeze
    headers = ["Name", "Address", "State", "Zip"]
    sheet["A1"] = headers[0]
    sheet["B1"] = headers[1]
    sheet["C1"] = headers[2]
    sheet["D1"] = headers[3]
    data = [dict(zip(headers, ("Mike", "123 Storm Dr", "IA", "50000"))), dict(zip(headers, ("Ted", "555 Tornado Alley", "OK", "90000")))]
    print(data)
    row = 2
    for d in data:
        sheet[f'A{row}'] = d["Name"]
        sheet[f'B{row}'] = d["Address"]
        sheet[f'C{row}'] = d["State"]
        sheet[f'D{row}'] = d["Zip"]
        row += 1
    wb.save(path)

if __name__ == "__main__":
    #data = {"B1": "Hi", "B5": "Python"}
    #edit("C:/Users/singhman/Downloads/learn_PyExcel/hello.xlsx", data)

    #merged_cells("C:/Users/singhman/Downloads/learn_PyExcel/cell_merged.xlsx", "Hello World")
    #folding("C:/Users/singhman/Downloads/learn_PyExcel/cell_fold.xlsx", rows=(1,5), cols=())
    freeze("C:/Users/singhman/Downloads/learn_PyExcel/freeze.xlsx", row_to_freeze="A2")
