import openpyxl
book = openpyxl.load_workbook("D:\\OneDrive - Mavenir Systems, Inc\\Documents\\pythondemo.xlsx")
sheet = book.active
Dict = {}
cell = sheet.cell(row=1, column=2)
print(cell.value)

sheet.cell(row=2, column=2).value = "Manmohan"
print(sheet.cell(row=2, column=2).value)

print(sheet.max_row)
print(sheet.max_column)

print(sheet['A3'].value) #shortcut to fetch value from a cell.

#for i in range(1, sheet.max_row+1):
#    if sheet.cell(row=i, column=1).value == "TestCase2": # get the data only from a row we want
#        for j in range(1, sheet.max_column+1):
#            print(sheet.cell(row=i, column=j).value)


for i in range(1, sheet.max_row+1):
    if sheet.cell(row=i, column=1).value == "TestCase2": # get the data only from a row we want
        for j in range(2, sheet.max_column+1):
            #Dict["firstname"] = "Manmohan"
            Dict[sheet.cell(row=1, column=j).value] = sheet.cell(row=i, column=j).value

print(Dict)
